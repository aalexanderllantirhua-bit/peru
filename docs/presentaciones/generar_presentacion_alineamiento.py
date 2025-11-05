"""Genera una presentación PPTX sobre el uso de equipos de alineamiento electrónico en vehículos."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


OUTPUT_PATH = Path(__file__).resolve().parent / "Resumen_Alineamiento_Electronico.pptx"


def add_title_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Alineamiento electrónico en vehículos"
    subtitle = slide.placeholders[1]
    subtitle.text = "Resumen de uso de equipos y mejores prácticas"


def add_objectives_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Objetivos de la presentación"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Comprender el propósito del alineamiento electrónico"
    for bullet in (
        "Identificar componentes principales del equipo",
        "Revisar flujo de trabajo típico",
        "Resaltar seguridad, mantenimiento y beneficios",
    ):
        p = body.add_paragraph()
        p.text = bullet
        p.level = 0


def add_equipment_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Descripción del equipo"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Equipo de alineamiento electrónico"
    details = (
        "Sensores y cámaras de alta precisión",
        "Plataformas de referencia y soportes",
        "Unidad de control con software especializado",
        "Interfaces de diagnóstico y conexión a la red",
    )
    for bullet in details:
        p = body.add_paragraph()
        p.text = bullet
        p.level = 0


def add_components_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[5]  # Title Only
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Componentes clave"

    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(4)
    height = Inches(4)

    text_box_left = slide.shapes.add_textbox(left, top, width, height)
    tf_left = text_box_left.text_frame
    tf_left.text = "Sensores y cámaras"
    for bullet in (
        "Capturan ángulos de convergencia y caída",
        "Detectan desviaciones con precisión milimétrica",
    ):
        p = tf_left.add_paragraph()
        p.text = bullet
        p.level = 1

    text_box_right = slide.shapes.add_textbox(left + Inches(4.2), top, width, height)
    tf_right = text_box_right.text_frame
    tf_right.text = "Software y hardware"
    for bullet in (
        "Procesa datos y genera informes",
        "Permite calibraciones asistidas paso a paso",
        "Interfaz gráfica intuitiva",
    ):
        p = tf_right.add_paragraph()
        p.text = bullet
        p.level = 1


def add_workflow_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Flujo de trabajo típico"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Etapas principales"
    steps = (
        "Inspección inicial del vehículo",
        "Selección del modelo y parámetros en el software",
        "Montaje de sensores y calibración base",
        "Medición y ajuste de ángulos",
        "Generación de informe y validación final",
    )
    for step in steps:
        p = body.add_paragraph()
        p.text = step
        p.level = 0


def add_calibration_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Calibración y verificación"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Buenas prácticas"
    bullet_points = (
        "Asegurar superficie nivelada y ruedas centradas",
        "Configurar referencias del fabricante",
        "Realizar comprobaciones cruzadas de lectura",
        "Registrar tolerancias y resultados para historial",
    )
    for bullet in bullet_points:
        p = body.add_paragraph()
        p.text = bullet
        p.level = 0


def add_safety_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Seguridad operacional"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Pautas esenciales"
    guidelines = (
        "Verificar bloqueo de ruedas y elevadores",
        "Evitar cables sueltos y obstáculos",
        "Utilizar EPP: guantes, calzado y gafas",
        "Actualizar firmware y software para correcciones",
    )
    for item in guidelines:
        p = body.add_paragraph()
        p.text = item
        p.level = 0


def add_maintenance_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Mantenimiento y buenas prácticas"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Rutinas recomendadas"
    tasks = (
        "Limpieza de lentes y sensores tras cada jornada",
        "Verificación periódica de cables y conectores",
        "Actualización de bases de datos de vehículos",
        "Documentar incidencias y soluciones aplicadas",
    )
    for task in tasks:
        p = body.add_paragraph()
        p.text = task
        p.level = 0


def add_benefits_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Beneficios del alineamiento electrónico"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Impacto en el servicio"
    benefits = (
        "Mayor precisión en ajustes y diagnósticos",
        "Reducción de tiempo de intervención",
        "Incremento de la satisfacción del cliente",
        "Mejor conservación de neumáticos y suspensión",
    )
    for benefit in benefits:
        p = body.add_paragraph()
        p.text = benefit
        p.level = 0


def add_conclusions_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Conclusiones"
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Puntos clave"
    conclusions = (
        "El alineamiento electrónico mejora la eficiencia del taller",
        "Requiere formación continua del personal",
        "Su éxito depende de mantenimiento y verificación constante",
        "Ofrece reportes claros para clientes y auditorías",
    )
    for conclusion in conclusions:
        p = body.add_paragraph()
        p.text = conclusion
        p.level = 0


def generate_presentation(output_path: Path = OUTPUT_PATH) -> None:
    prs = Presentation()
    add_title_slide(prs)
    add_objectives_slide(prs)
    add_equipment_slide(prs)
    add_components_slide(prs)
    add_workflow_slide(prs)
    add_calibration_slide(prs)
    add_safety_slide(prs)
    add_maintenance_slide(prs)
    add_benefits_slide(prs)
    add_conclusions_slide(prs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(f"Presentación generada en: {output_path}")


if __name__ == "__main__":
    try:
        generate_presentation()
    except Exception as exc:  # pragma: no cover
        print(f"Error al generar la presentación: {exc}")
